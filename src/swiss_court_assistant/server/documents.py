"""What a client hands over, turned into text the assistant can work on.

A matter starts either from a document or from a recording of the client telling their story.
Documents are parsed and stored by parsing.py (Nemotron Parse for PDFs and scans); this module keeps
the Word reader it uses and the transcription of recordings. Recordings arrive as 16 kHz mono PCM
because the browser decodes and resamples them: every format the browser can play then works without
ffmpeg on the server, and the bytes are already exactly what the ASR NIM expects.
"""

from __future__ import annotations

import asyncio
import io
import logging
import re

from .voice import SAMPLE_RATE, Listener

log = logging.getLogger(__name__)

# About twenty pages of a matter's document. Intake only needs the story, and the whole text goes
# into one prompt.
MAX_CHARS = 40_000
CHUNK = SAMPLE_RATE * 2 // 5  # 200 ms of 16-bit mono audio
_BLANK_LINES = re.compile(r"\n{3,}")
_TRAILING = re.compile(r"[ \t]+$", re.M)


class UnreadableError(Exception):
    """The file holds no text we can use - an empty or damaged file, or a format we do not read."""


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # termination letters and settlement offers often are tables
        lines += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
    return "\n".join(lines)


def _tidy(text: str) -> str:
    return _BLANK_LINES.sub("\n\n", _TRAILING.sub("", text.replace("\f", "\n"))).strip()


async def transcribe(listener: Listener, pcm: bytes, language: str = "en") -> str:
    """A recording (16 kHz mono PCM) as text, through the same ASR NIM as voice mode."""
    audio: asyncio.Queue[bytes | None] = asyncio.Queue()
    for i in range(0, len(pcm), CHUNK):
        audio.put_nowait(pcm[i:i + CHUNK])
    audio.put_nowait(None)
    said = [text async for text, final in listener.transcribe(audio, language) if final and text]
    out = _tidy(" ".join(said))
    if not out:
        raise UnreadableError("Nothing was recognised in that recording.")
    return out[:MAX_CHARS]
