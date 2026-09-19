"""A matter: one client case carried through the stages software can actually help with.

A lawyer's case moves through five stages. Four of them are here:

1. **Intake** — the client tells a messy story; the legal questions in it have to be spotted.
2. **Research** — for each question: which provision applies, what do the courts say, is that still
   the current line. This is the existing agent, run once per issue.
3. **Assessment** — how strong is the position, and what will the other side argue.
4. **Drafting** — a memo in which every proposition carries a citation.

The fifth, filing and deadlines and client files, is practice management: no case law would help,
so the app says so rather than pretending.

Nothing here invents legal content: the issues and the assessment are written from the client's own
text and from what the research actually retrieved, and the memo is assembled from the researched
answers with their citations kept intact.
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import re
import sqlite3
import threading
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from langchain_core.messages import HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI
from pydantic import BaseModel, ValidationError

from .agent import Agent, Cite, Delta, ToolStart, Verdict
from .case_index import CaseIndex, IndexUnavailable, collection_of
from .language import detect_language
from .mentions import statute_links
from .llm import LLM_KEY, LLM_URL, served_model
from .schemas import DocumentInfo, Intake, Issue, Matter, MatterSummary, Source
from .store import new_id, now

log = logging.getLogger(__name__)

MAX_ISSUES = 4  # each one is a full research run; four already takes minutes
LANGUAGE_NAMES = {"de": "German", "fr": "French", "it": "Italian", "rm": "Romansh", "en": "English"}


# ── intake ──────────────────────────────────────────────────────────────
class _IssueDraft(BaseModel):
    question: str
    why: str
    area: str


class _Intake(BaseModel):
    title: str
    summary: str
    parties: list[str]
    timeline: list[str]
    issues: list[_IssueDraft]


INTAKE_FORMAT = {"type": "json_schema",
                 "json_schema": {"name": "intake", "schema": _Intake.model_json_schema()}}

INTAKE_PROMPT = """You are a Swiss lawyer taking in a new matter. You are given what the client handed over — a document, or a transcript of the client telling their story. Read it and return JSON:

- "title": a short name for the matter, e.g. "Termination of a lease in Zurich" (in {language}).
- "summary": three to five sentences stating the facts as they are given, in {language}. Facts only, no legal conclusions, no advice.
- "parties": who is involved, one entry each, with their role ("Client — tenant", "Landlord (AG)").
- "timeline": the dated facts in order, one per entry ("2024-03-14: termination served"). Leave empty if the text gives no dates.
- "issues": between one and {max_issues} legal questions that decide this matter, most important first. Each has "question": the question as a lawyer would research it in Swiss case law, in {language}; "why": one sentence on why it decides this case; "area": one of "Zivilrecht", "Strafrecht", "öffentliches Recht", "Sozialversicherungsrecht".

Name a statute article in a question **only** if the client's text names it. Writing an article number from memory puts a wrong provision into the research — ask the question in words instead ("termination during pregnancy", not "art. 336c CO").

Only use what the text says. If a fact is missing, do not invent it — an issue may note that it is open."""

# A statute reference in an intake question, with whatever leads into it ("nach Art. 271a Abs. 1
# lit. d OR", "selon l'art. 336c CO").
_ARTICLE = re.compile(
    r"(?:\s*(?:nach|gem(?:ä|ae)ss|i\.S\.v\.|im Sinne von|selon|au sens de|conform(?:é|e)ment (?:à|a)|"
    r"secondo|ai sensi di|under|according to|pursuant to|et|und|and|sowie|oder|ou)\s*"
    r"(?:l'|la|le|les|der|dem|den|des|dell'|della|the)?\s*)?"
    r"\bArt(?:\.|icle|ikel)?\s*(\d+[a-z]?)"
    r"(?:\s*(?:Abs|al|cpv|para)\.?\s*\d+)?(?:\s*(?:lit|let|lett)\.?\s*[a-z])?"
    r"(?:\s+(?:[A-Z][A-Za-z]{1,5}))?", re.I)


def without_invented_articles(question: str, facts: str) -> str:
    """Drop statute references the client's own text never mentions.

    Asked for the question a lawyer would research, the model likes to name an article from memory,
    and a wrong one (seen: "Art. 269 OR" for a termination, "art. 12 CO" for overtime) sends the
    research after the wrong provision. An article number that appears in the facts is the client's
    and stays; anything else is dropped and the question stands in words."""
    def keep(m: re.Match[str]) -> str:
        return m.group(0) if re.search(rf"\b{re.escape(m.group(1))}\b", facts) else ""

    cleaned = re.sub(r"[ \t]{2,}", " ", re.sub(r"\s+([,.])", r"\1", _ARTICLE.sub(keep, question)))
    cleaned = re.sub(r"[,:;]\s*(\?|$)", r"\1", cleaned.strip(" ,;:-–"))
    return cleaned or question


# ── assessment ──────────────────────────────────────────────────────────
ASSESS_PROMPT = """You are a Swiss lawyer assessing a matter for the file, writing in {language}.

You are given the facts and, for each legal issue, what case-law research found. Write a short assessment in Markdown with these sections:

**Where the client stands** — two or three sentences.
**In our favour** — bullets. Each bullet must rest on something the research found; keep the [n] citation markers from the research text exactly as they are, so the reader can check it.
**Against us** — bullets: what the other side will argue. Where a bullet rests on a decision the research found — one that cuts the other way, or is silent — keep its [n] marker too, exactly as it appears in the research text; only a bullet that states no case law was found may go without one.
**What is still open** — bullets: facts or documents needed before this can be advised on; these describe a gap, so they carry no [n].

Write the four headings in {language} as well, in bold — not in English.

Rules: use only the facts given and the research results below. Every claim about what a decision holds must carry the [n] that supports it — never state what the case law says without its marker. Where the research found nothing on point, say that plainly instead of filling the gap from your own knowledge. Do not predict a percentage chance of success. Do not give the client instructions; this is a note for the lawyer's file."""


def _digest(matter: Matter) -> str:
    """The researched issues as one block for the assessment prompt, with the citation markers already
    renumbered the way the memo numbers them — otherwise every issue contributes its own [1] and the
    assessment's markers point at the wrong decisions."""
    _, answers = numbered(matter)
    out = [f"FACTS:\n{matter.intake.summary if matter.intake else matter.facts[:2000]}"]
    for issue, answer in zip(matter.issues, answers, strict=True):
        found = answer.strip() or "Nothing on point was found."
        out.append(f"\nISSUE {issue.n}: {issue.question}\nWhy it matters: {issue.why}\nRESEARCH:\n{found}")
    return "\n".join(out)


# ── the memo ────────────────────────────────────────────────────────────
HEADINGS = {
    "de": ("Aktennotiz", "Sachverhalt", "Beteiligte", "Chronologie", "Rechtsfragen", "Einschätzung",
           "Quellenverzeichnis", "Nicht abgedeckt"),
    "fr": ("Note de dossier", "Faits", "Parties", "Chronologie", "Questions juridiques", "Appréciation",
           "Table des autorités", "Non couvert"),
    "it": ("Nota interna", "Fatti", "Parti", "Cronologia", "Questioni giuridiche", "Valutazione",
           "Indice delle fonti", "Non coperto"),
    "en": ("Memorandum", "Facts", "Parties", "Timeline", "Issues", "Assessment",
           "Table of authorities", "Not covered"),
}
DISCLAIMER = {
    "de": "Erstellt mit dem Swiss Court Assistant aus Schweizer Gerichtsentscheiden. Jede Fundstelle ist "
          "zu prüfen; Fristen, Verfahrensstand und Lehre sind nicht abgedeckt.",
    "fr": "Établi avec le Swiss Court Assistant à partir de décisions suisses. Chaque référence doit être "
          "vérifiée ; délais, état de la procédure et doctrine ne sont pas couverts.",
    "it": "Redatto con lo Swiss Court Assistant sulla base di decisioni svizzere. Ogni riferimento va "
          "verificato; termini, stato della procedura e dottrina non sono coperti.",
    "en": "Drafted with the Swiss Court Assistant from Swiss court decisions. Every reference has to be "
          "checked; deadlines, the state of the proceedings and legal doctrine are not covered.",
}


def _renumber(text: str, mapping: dict[int, int]) -> str:
    """An issue's local [1], [2] … as the memo's running numbers."""
    for local in sorted(mapping, reverse=True):  # high first, so [1]->[10] is not renumbered again
        text = text.replace(f"[{local}]", f"[[{mapping[local]}]]")
    return text.replace("[[", "[").replace("]]", "]")


def numbered(matter: Matter) -> tuple[list[Source], list[str]]:
    """Every issue's citations on one running series: the memo's authorities, and each issue's answer
    with its local [1], [2] … rewritten to that series."""
    sources: list[Source] = []
    answers: list[str] = []
    for issue in matter.issues:
        mapping = {}
        for source in issue.sources or []:
            sources.append(source)
            mapping[source.n] = len(sources)
        answers.append(_renumber(issue.answer or "", mapping))
    return sources, answers


def case_brief(matter: Matter) -> str:
    """The matter's case prep as background for a question asked about it in the assistant.

    Each citation marker becomes the decision id behind it, so the agent knows which decisions to read
    again — the brief is not a source itself: a statement the answer makes still needs a passage the
    agent fetched in its own turn, or the grounding check drops it."""
    sources, answers = numbered(matter)

    def leads(text: str) -> str:
        def one(m: re.Match) -> str:
            n = int(m.group(1))
            return f" (see {sources[n - 1].decision_id})" if 0 < n <= len(sources) else ""
        return re.sub(r"\s*\[(\d+)\]", one, text).strip()

    out = [f"CASE PREP — the matter the user is asking about: {matter.title}",
           "This is the lawyer's own preparation of the case, so you know what it is about. It is background, "
           "not a source: to state what a decision, a statute or the case file says, read it with your tools in "
           "this turn and cite that. The decision ids below are where the earlier research found each point.",
           f"\nFACTS:\n{matter.intake.summary if matter.intake else matter.facts[:2000]}"]
    if matter.intake and matter.intake.parties:
        out.append("PARTIES: " + "; ".join(matter.intake.parties))
    if matter.intake and matter.intake.timeline:
        out.append("TIMELINE:\n" + "\n".join(f"- {t}" for t in matter.intake.timeline))
    for issue, answer in zip(matter.issues, answers, strict=True):
        found = leads(answer)[:2500] or "Not researched yet."
        out.append(f"\nISSUE {issue.n}: {issue.question}\nWhy it matters: {issue.why}\nRESEARCH:\n{found}")
    if matter.assessment:
        out.append(f"\nASSESSMENT:\n{leads(matter.assessment)[:3000]}")
    return "\n".join(out)


def memo(matter: Matter) -> str:
    """The matter as one Markdown memo, assembled from the researched answers — not rewritten by a
    model, so the citations stay attached to the sentences the research put them on."""
    lang = matter.language if matter.language in HEADINGS else "en"
    head, facts, parties, timeline, issues, assessment, authorities, uncovered = HEADINGS[lang]
    out = [f"# {head}: {matter.title}", "", f"*{matter.created_at[:10]} · {matter.id}*", "",
           f"## {facts}", "", (matter.intake.summary if matter.intake else matter.facts[:1500]), ""]
    if matter.intake and matter.intake.parties:
        out += [f"## {parties}", "", *[f"- {p}" for p in matter.intake.parties], ""]
    if matter.intake and matter.intake.timeline:
        out += [f"## {timeline}", "", *[f"- {t}" for t in matter.intake.timeline], ""]

    sources, answers = numbered(matter)
    out += [f"## {issues}", ""]
    for issue, answer in zip(matter.issues, answers, strict=True):
        out += [f"### {issue.n}. {issue.question}", "", f"*{issue.why}*", "", answer or "_—_", ""]
    if matter.assessment:
        out += [f"## {assessment}", "", matter.assessment, ""]
    if sources:
        out += [f"## {authorities}", ""]
        for i, s in enumerate(sources, 1):
            d = s.decision
            # one entry per cited passage, not per decision, so the considerandum tells two passages
            # of the same decision apart
            erw = (", ".join(s.erwaegungen) if s.section == "document" else f"E. {', '.join(s.erwaegungen)}") \
                if s.erwaegungen else None  # a page of the client's document: "p. 2"
            where = " · ".join(x for x in (d.court_label, d.docket, erw, d.date) if x)
            link = f" — {d.source_url}" if d.source_url else ""
            mark = "" if s.supported is not False else "  ⚠ check: the passage may not state this"
            out.append(f"{i}. {where}{link}{mark}")
        out.append("")
    out += [f"## {uncovered}", "",
            "- " + DISCLAIMER[lang], ""]
    return "\n".join(out)


# ── the memo as a Word file ─────────────────────────────────────────────
_BULLET = re.compile(r"^\s*[-*•]\s+")
_INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|_(.+?)_")


def _runs(paragraph, text: str) -> None:
    """Markdown emphasis as Word runs: **bold**, *italic*, _italic_."""
    at = 0
    for m in _INLINE.finditer(text):
        if m.start() > at:
            paragraph.add_run(text[at:m.start()])
        run = paragraph.add_run(m.group(1) or m.group(2) or m.group(3) or "")
        run.bold = m.group(1) is not None
        run.italic = m.group(1) is None
        at = m.end()
    if at < len(text):
        paragraph.add_run(text[at:])


def _body(document, text: str) -> None:
    """A researched answer or the assessment: its paragraphs, bullets and emphasis into the document.
    The [n] markers are left exactly as they are — they point into the table of authorities."""
    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):  # a heading inside an answer becomes a bold line, not a Word heading
            paragraph = document.add_paragraph()
            paragraph.add_run(stripped.lstrip("# ").strip()).bold = True
        elif _BULLET.match(stripped):
            _runs(document.add_paragraph(style="List Bullet"), _BULLET.sub("", stripped))
        else:
            _runs(document.add_paragraph(), stripped)


def docx_memo(matter: Matter) -> bytes:
    """The same memo as a Word file, so it can go straight into the client file.

    Built from the matter, not from the Markdown, for the same reason `memo()` is assembled rather
    than written by a model: the citations stay on the sentences the research put them on."""
    import docx
    from docx.shared import Pt

    lang = matter.language if matter.language in HEADINGS else "en"
    head, facts, parties, timeline, issues, assessment, authorities, uncovered = HEADINGS[lang]
    document = docx.Document()
    document.add_heading(f"{head}: {matter.title}", level=0)
    stamp = document.add_paragraph()
    run = stamp.add_run(f"{matter.created_at[:10]} · {matter.id}")
    run.italic = True
    run.font.size = Pt(9)

    document.add_heading(facts, level=1)
    document.add_paragraph(matter.intake.summary if matter.intake else matter.facts[:1500])
    if matter.intake and matter.intake.parties:
        document.add_heading(parties, level=1)
        for party in matter.intake.parties:
            document.add_paragraph(party, style="List Bullet")
    if matter.intake and matter.intake.timeline:
        document.add_heading(timeline, level=1)
        for entry in matter.intake.timeline:
            document.add_paragraph(entry, style="List Bullet")

    sources, answers = numbered(matter)
    if matter.issues:
        document.add_heading(issues, level=1)
        for issue, answer in zip(matter.issues, answers, strict=True):
            document.add_heading(f"{issue.n}. {issue.question}", level=2)
            document.add_paragraph().add_run(issue.why).italic = True
            _body(document, answer or "—")
    if matter.assessment:
        document.add_heading(assessment, level=1)
        _body(document, matter.assessment)
    if sources:
        document.add_heading(authorities, level=1)
        for i, source in enumerate(sources, 1):
            d = source.decision
            erw = f"E. {', '.join(source.erwaegungen)}" if source.erwaegungen else None
            where = " · ".join(x for x in (d.court_label, d.docket, erw, d.date) if x)
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(f"{i}. {where}")
            if d.source_url:
                link = paragraph.add_run(f"\n{d.source_url}")
                link.font.size = Pt(8)
            if source.supported is False:  # what the grounding check flagged, carried into the file
                warning = paragraph.add_run("\n⚠ check: the passage may not state this")
                warning.bold = True
    document.add_heading(uncovered, level=1)
    document.add_paragraph(DISCLAIMER[lang], style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── storage ─────────────────────────────────────────────────────────────
_SCHEMA = """
CREATE TABLE IF NOT EXISTS matters (
    id TEXT PRIMARY KEY,
    title TEXT NOT NULL,
    stage TEXT NOT NULL,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    data TEXT NOT NULL  -- JSON Matter
);
"""


class MatterStore:
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self._db = sqlite3.connect(path, check_same_thread=False)
        self._db.row_factory = sqlite3.Row
        self._db.executescript(_SCHEMA)
        self._lock = threading.Lock()

    def list(self) -> list[MatterSummary]:
        with self._lock:
            rows = self._db.execute("SELECT data FROM matters ORDER BY updated_at DESC").fetchall()
        return [MatterSummary(**{k: v for k, v in json.loads(r["data"]).items()
                                 if k in MatterSummary.model_fields}) for r in rows]

    def get(self, matter_id: str) -> Matter | None:
        with self._lock:
            row = self._db.execute("SELECT data FROM matters WHERE id = ?", (matter_id,)).fetchone()
        return Matter.model_validate_json(row["data"]) if row else None

    def save(self, matter: Matter) -> Matter:
        matter.updated_at = now()
        with self._lock, self._db:
            self._db.execute(
                "INSERT INTO matters (id, title, stage, created_at, updated_at, data) VALUES (?,?,?,?,?,?) "
                "ON CONFLICT(id) DO UPDATE SET title=excluded.title, stage=excluded.stage, "
                "updated_at=excluded.updated_at, data=excluded.data",
                (matter.id, matter.title, matter.stage, matter.created_at, matter.updated_at,
                 matter.model_dump_json()))
        return matter

    def delete(self, matter_id: str) -> bool:
        with self._lock, self._db:
            return self._db.execute("DELETE FROM matters WHERE id = ?", (matter_id,)).rowcount > 0


def create_matter(store: MatterStore, facts: str, source_name: str | None, source_kind: str,
           title: str | None = None, assets: list[DocumentInfo] | None = None) -> Matter:
    assets = assets or []
    stamp = datetime.now(UTC).isoformat(timespec="milliseconds")
    matter = Matter(id=new_id(), title=title or (source_name or facts[:60].strip() or "New matter"),
                    stage="new", source_name=source_name, source_kind=source_kind,  # type: ignore[arg-type]
                    document_id=assets[0].id if assets else None, assets=assets,
                    language=detect_language(facts), facts=facts, created_at=stamp, updated_at=stamp)
    return store.save(matter)


# ── the pipeline ────────────────────────────────────────────────────────
class Pipeline:
    """Runs a matter through the four stages, saving after each one and streaming what it does."""

    def __init__(self, agent: Agent, store: MatterStore, decisions: Any = None, max_issues: int = MAX_ISSUES,
                 case_index: CaseIndex | None = None):
        self.agent, self.store, self.decisions, self.max_issues = agent, store, decisions, max_issues
        self.case_index = case_index
        common: dict[str, Any] = dict(base_url=LLM_URL, api_key=LLM_KEY, model=served_model(),
                                      temperature=0.2, streaming=True)
        self.intake_llm = ChatOpenAI(**common, max_tokens=2048, extra_body={
            "chat_template_kwargs": {"enable_thinking": False}, "response_format": INTAKE_FORMAT})
        self.assess_llm = ChatOpenAI(**common, max_tokens=1500, extra_body={
            "chat_template_kwargs": {"enable_thinking": False}})

    async def _intake(self, matter: Matter) -> AsyncIterator[dict[str, Any]]:
        language = LANGUAGE_NAMES.get(matter.language, "English")
        prompt = INTAKE_PROMPT.format(language=language, max_issues=self.max_issues)
        reply = await self.intake_llm.ainvoke([SystemMessage(prompt), HumanMessage(matter.facts)])
        try:
            drafted = _Intake.model_validate_json(str(reply.content))
        except (ValidationError, ValueError) as e:
            raise RuntimeError("The intake step did not return a readable result.") from e
        matter.title = drafted.title.strip() or matter.title
        matter.intake = Intake(summary=drafted.summary, parties=drafted.parties, timeline=drafted.timeline)
        matter.issues = [
            Issue(n=i, question=without_invented_articles(d.question, matter.facts), why=d.why, area=d.area)
            for i, d in enumerate(drafted.issues[:self.max_issues], 1)]
        yield {"type": "intake", "title": matter.title,
               "intake": matter.intake.model_dump(by_alias=True),
               "issues": [i.model_dump(by_alias=True) for i in matter.issues]}

    async def _research(self, matter: Matter, issue: Issue) -> AsyncIterator[dict[str, Any]]:
        """One issue through the ordinary research agent, so it gets the same searches, citations
        and grounding checks as a question asked in the chat."""
        context = matter.intake.summary if matter.intake else matter.facts[:800]
        question = f"{issue.question}\n\n(Context — the matter this is asked for: {context})"
        parts: list[str] = []
        sources: list[Source] = []
        yield {"type": "issue_start", "n": issue.n}
        # The client's own document, so the research can quote it for the facts. Without it the answer
        # stated the facts in a sentence cited to a court decision, the check found the decision did not
        # say them, and dropped the sentence together with the law in it.
        documents = getattr(self.agent, "documents", None)
        attached = list(matter.assets)
        if not attached and matter.document_id and documents and (info := documents.info(matter.document_id)):
            attached = [info]  # a matter opened before the case file was kept
        # a long case file is searched in its collection rather than shown by its documents' beginnings
        collection = collection_of(matter.id) if matter.indexed else None
        async for ev in self.agent.answer(question, [], ask=False, attachments=attached,  # no one to ask mid-memo
                                          collection=collection):
            match ev:
                case ToolStart():
                    arg = " ".join(str(v) for v in ev.args.values() if v not in (None, "", False))
                    yield {"type": "issue_tool", "n": issue.n, "name": ev.name, "arg": arg[:120]}
                case Delta():
                    parts.append(ev.text)
                    yield {"type": "issue_delta", "n": issue.n, "text": ev.text}
                case Cite():
                    if all(s.n != ev.source.n for s in sources):
                        sources.append(ev.source)
                    parts.append(f"[{ev.source.n}]")
                    yield {"type": "issue_citation", "n": issue.n,
                           "source": ev.source.model_dump(by_alias=True)}
                case Verdict():
                    for source in sources:
                        if source.n == ev.n:
                            source.supported = ev.supported
                    yield {"type": "issue_verdict", "n": issue.n, "source": ev.n, "supported": ev.supported}
        issue.answer, issue.sources = "".join(parts), sources
        if self.decisions is not None:
            issue.statutes = await asyncio.to_thread(statute_links, self.decisions, issue.answer, matter.language) or None
        yield {"type": "issue_done", "n": issue.n, "issue": issue.model_dump(by_alias=True)}

    async def _assess(self, matter: Matter) -> AsyncIterator[dict[str, Any]]:
        prompt = ASSESS_PROMPT.format(language=LANGUAGE_NAMES.get(matter.language, "English"))
        parts: list[str] = []
        async for chunk in self.assess_llm.astream([SystemMessage(prompt), HumanMessage(_digest(matter))]):
            text = str(chunk.content)
            if text:
                parts.append(text)
                yield {"type": "assessment_delta", "text": text}
        matter.assessment = "".join(parts).strip()

    async def _index(self, matter: Matter) -> AsyncIterator[dict[str, Any]]:
        """Put the case file into its own collection of the case index: cut into passages and embedded, so
        each issue's research can search all of it. Idempotent — a rerun only adds what is missing. Without
        the embedder the research falls back to the documents' beginnings and read_document."""
        if self.case_index is None or not matter.assets:
            return
        collection = collection_of(matter.id)
        try:
            await asyncio.to_thread(self.case_index.add, collection, matter.assets)
        except IndexUnavailable as e:
            log.warning("matter %s: case file not indexed (%s)", matter.id, e)
            return
        matter.indexed = await asyncio.to_thread(self.case_index.count, collection)
        yield {"type": "indexed", "passages": matter.indexed}

    async def run(self, matter: Matter) -> AsyncIterator[dict[str, Any]]:
        """Stage by stage, saving as it goes: a browser that disconnects loses the stream, not the work."""
        try:
            # a rerun starts over: the last run's assessment and memo belong to research it replaces
            matter.stage, matter.assessment, matter.memo = "intake", None, None
            yield {"type": "stage", "stage": "intake", "status": "running"}
            async for ev in self._index(matter):
                yield ev
            async for ev in self._intake(matter):
                yield ev
            self.store.save(matter)
            yield {"type": "stage", "stage": "intake", "status": "done"}

            matter.stage = "research"
            yield {"type": "stage", "stage": "research", "status": "running"}
            for issue in matter.issues:
                async for ev in self._research(matter, issue):
                    yield ev
                self.store.save(matter)
            yield {"type": "stage", "stage": "research", "status": "done"}

            matter.stage = "assessment"
            yield {"type": "stage", "stage": "assessment", "status": "running"}
            async for ev in self._assess(matter):
                yield ev
            self.store.save(matter)
            yield {"type": "stage", "stage": "assessment", "status": "done"}

            matter.stage = "drafting"
            yield {"type": "stage", "stage": "drafting", "status": "running"}
            matter.memo = memo(matter)
            matter.stage = "done"
            self.store.save(matter)
            yield {"type": "stage", "stage": "drafting", "status": "done"}
            yield {"type": "done", "matter": matter.model_dump(by_alias=True)}
        except Exception as e:
            log.exception("matter %s failed in stage %s", matter.id, matter.stage)
            self.store.save(matter)
            yield {"type": "error", "stage": matter.stage, "message": str(e) or "The step failed."}
